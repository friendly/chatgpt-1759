"""Build reference.docx (Word style sheet) for the Statistique et Societe submission.

Follows the author instructions (2024-01-17): Times New Roman 12, line spacing 1.2,
12 pt before/after each paragraph, numbered headings, footnotes, styles for long
quotations, and illustration title / legend / credits.  The first-page (front matter)
styles imitate a real submission made on the journal's Word template: Calibri title
block, then Courier New 10 pt navy text indented 1 cm.

Once the journal's own French .docx template is available, replace reference.docx by
it (or map its style names onto the custom-style names used in chronique.qmd).
"""
import os
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
BASE = ROOT / "_base_reference.docx"
with open(BASE, "wb") as fh:
    fh.write(subprocess.check_output(["quarto", "pandoc", "--print-default-data-file", "reference.docx"]))
doc = Document(BASE)
S = doc.styles

J = WD_ALIGN_PARAGRAPH.JUSTIFY
L = WD_ALIGN_PARAGRAPH.LEFT
C = WD_ALIGN_PARAGRAPH.CENTER
TNR = "Times New Roman"


def set_font(style, name, size=None, bold=None, italic=None, color=None):
    """Set the font of a style, dropping theme fonts/colours inherited from pandoc's sheet."""
    rpr = style.element.get_or_add_rPr()
    for el in rpr.findall(qn("w:rFonts")):
        rpr.remove(el)
    rf = OxmlElement("w:rFonts")
    for a in ("ascii", "hAnsi", "cs", "eastAsia"):
        rf.set(qn("w:" + a), name)
    rpr.insert(0, rf)
    for el in rpr.findall(qn("w:color")):
        rpr.remove(el)
    f = style.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = RGBColor.from_string(color)


def para(style, before=12, after=12, line=1.2, align=None, left=None, right=None, first=None,
         keep_next=None, keep_lines=None):
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        pf.alignment = align
    if left is not None:
        pf.left_indent = Cm(left)
    if right is not None:
        pf.right_indent = Cm(right)
    if first is not None:
        pf.first_line_indent = Cm(first)
    if keep_next is not None:
        pf.keep_with_next = keep_next
    if keep_lines is not None:
        pf.keep_together = keep_lines
    ppr = style.element.get_or_add_pPr()
    for tag in ("w:contextualSpacing",):
        for el in ppr.findall(qn(tag)):
            ppr.remove(el)


def get(name, kind=WD_STYLE_TYPE.PARAGRAPH, base="Normal"):
    try:
        return S[name]
    except KeyError:
        st = S.add_style(name, kind)
        if kind == WD_STYLE_TYPE.PARAGRAPH:
            st.base_style = S[base]
        return st


def strip_ppr(style, *tags):
    ppr = style.element.get_or_add_pPr()
    for t in tags:
        for el in ppr.findall(qn(t)):
            ppr.remove(el)


# ---- body ----------------------------------------------------------------------------------
for n in ("Normal", "Body Text", "First Paragraph"):
    st = get(n)
    set_font(st, TNR, 12, bold=False, italic=False, color="000000")
    para(st, 12, 12, 1.2, J)
st = get("Compact")
set_font(st, TNR, 10, color="000000")
para(st, 2, 2, 1.0, L)

# ---- headings (numbered "1", "1.1", "1.1.1" by Quarto number-sections) ---------------------
for n, sz in (("Heading 1", 18), ("Heading 2", 16), ("Heading 3", 14)):
    st = get(n)
    set_font(st, TNR, sz, bold=False, italic=False, color="000000")
    para(st, 12, 12, 1.2, L, keep_next=True, keep_lines=True)

# ---- footnotes -----------------------------------------------------------------------------
st = get("Footnote Text")
set_font(st, TNR, 10, color="000000")
para(st, 0, 3, 1.0, J)
try:
    S["Footnote Reference"].font.superscript = True
except KeyError:
    pass

# ---- long quotation ------------------------------------------------------------------------
st = get("Block Text")
set_font(st, TNR, 11, italic=False, color="000000")
para(st, 6, 6, 1.2, J, left=1.0, right=1.0)
strip_ppr(st, "w:pBdr", "w:shd")

# ---- code / worked examples (Courier) ------------------------------------------------------
st = get("Source Code")
set_font(st, "Courier New", 9.5, color="000000")
para(st, 6, 6, 1.0, L, left=1.0)
strip_ppr(st, "w:shd")
try:
    set_font(S["Verbatim Char"], "Courier New", 10, color="000000")
except KeyError:
    pass

# ---- illustration title / legend / credits (grey Arial, indented, as in the model) ---------
for name, sz, before, after in (("Titre illustration", 11, 12, 6),
                                ("Légende illustration", 9, 3, 3),
                                ("Crédits illustration", 9, 3, 12)):
    st = get(name)
    set_font(st, "Arial", sz, color="7F7F7F")
    para(st, before, after, 1.0, L, left=1.0)
S["Titre illustration"].paragraph_format.keep_with_next = True
S["Légende illustration"].paragraph_format.keep_with_next = True
for name in ("Table Caption", "Image Caption"):
    st = get(name)
    set_font(st, "Arial", 11, italic=False, color="7F7F7F")
    para(st, 12, 6, 1.0, L, left=1.0)
for name in ("Figure", "Captioned Figure"):
    st = get(name)
    set_font(st, TNR, 12, color="000000")
    para(st, 6, 6, 1.0, C, keep_next=True)

# ---- bibliography --------------------------------------------------------------------------
st = get("Bibliography")
set_font(st, TNR, 12, color="000000")
para(st, 0, 12, 1.2, J, left=1.0, first=-1.0)

# ---- hyperlinks: no underline (journal rule), discreet blue --------------------------------
try:
    h = S["Hyperlink"]
    h.font.underline = False
    h.font.color.rgb = RGBColor.from_string("1F3E99")
except KeyError:
    pass

# ---- front-matter styles (page 1) ----------------------------------------------------------
st = get("Titre article")
set_font(st, "Calibri Light", 28, color="000000")
para(st, 0, 6, 1.0, L)
st = get("Sous-titre article")
set_font(st, "Calibri", 11, color="5A5A5A")
para(st, 6, 12, 1.0, L)
st = get("Auteur")
set_font(st, "Calibri", 11, color="000000")
para(st, 6, 0, 1.0, L)
st = get("Front matter")
set_font(st, "Courier New", 10, color="000080")
para(st, 6, 6, 1.0, J, left=1.0, right=1.0)

# ---- tables: thick outer frame, thin inner rules (model in the instructions) ---------------
try:
    tbl = S["Table"]
    tblPr = tbl.element.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.element.append(tblPr)
    for el in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(el)
    b = OxmlElement("w:tblBorders")
    for side, sz in (("top", 12), ("left", 12), ("bottom", 12), ("right", 12),
                     ("insideH", 4), ("insideV", 4)):
        e = OxmlElement("w:" + side)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        b.append(e)
    tblPr.append(b)
except KeyError:
    pass

# ---- page: A4, 2.5 cm margins, page number bottom right ------------------------------------
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.5)
ft = sec.footer
ft.is_linked_to_previous = False
p = ft.paragraphs[0] if ft.paragraphs else ft.add_paragraph()
for r in list(p.runs):
    r._r.getparent().remove(r._r)
p.alignment = C if False else WD_ALIGN_PARAGRAPH.RIGHT
run = p.add_run()
for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
    if kind:
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), kind)
        run._r.append(fc)
    else:
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = text
        run._r.append(it)
run.font.name = TNR
run.font.size = Pt(12)

# empty body: pandoc only takes styles, page setup and footers from a reference document
body = doc.element.body
for el in list(body):
    if el.tag != qn("w:sectPr"):
        body.remove(el)
doc.save(ROOT / "reference.docx")
os.remove(BASE)
print("reference.docx written")
