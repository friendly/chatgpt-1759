"""Build the two Word versions of the chronique and check the journal's limits.

    python build.py          # render both .docx
    python build.py --pdf    # also export PDFs through Microsoft Word (for visual checks)

Outputs
    chronique_bibliography.docx  author-date citations + Bibliography section (journal standard, APA 7)
    chronique_footnotes.docx     every source as a footnote (full reference at first citation, short after)
"""
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT

HERE = Path(__file__).parent
SRC = HERE / "chronique.qmd"
text = SRC.read_text(encoding="utf-8")

BIB_BLOCK = re.compile(r"<!--BIBLIOGRAPHY-->.*?<!--/BIBLIOGRAPHY-->\n?", re.S)
MARKERS = re.compile(r"<!--/?BIBLIOGRAPHY-->\n?")

variants = {
    "bibliography": MARKERS.sub("", text),
    "footnotes": BIB_BLOCK.sub("", text).replace(
        "csl: assets/apa7.csl", "csl: assets/apa7-notes.csl\nfilters:\n  - assets/nosuppress.lua"),
}

def postprocess(path):
    """Page break before the Introduction (article text starts on page 2) and heading numbers without a dot."""
    d = Document(str(path))
    first = True
    for p in d.paragraphs:
        if p.style.name.startswith("Heading"):
            if p.runs:
                p.runs[0].text = re.sub(r"^(\d+)\. ", r"\g<1> ", p.runs[0].text)
            if first:
                p.paragraph_format.page_break_before = True
                first = False
    # keep each table on one page, together with its reading note and credits
    for tbl in d.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True
    d.save(str(path))


# Render outside Dropbox: Quarto's clean-up of its temporary folder fails on Dropbox's file locks.
with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmpdir:
    work = Path(tmpdir)
    for f in ("references.yaml", "reference.docx", "Annex.qmd"):
        shutil.copy(HERE / f, work / f)
    shutil.copytree(HERE / "assets", work / "assets")  # apa7.csl, apa7-notes.csl, nosuppress.lua
    shutil.copy(HERE / "R" / "oracle.R", work / "oracle.R")  # Annex.qmd sources it as a flat sibling
    shutil.copytree(HERE / "images", work / "images")
    for name, body in variants.items():
        (work / f"_build_{name}.qmd").write_text(body, encoding="utf-8")
        subprocess.run(["quarto", "render", f"_build_{name}.qmd", "--to", "docx"], cwd=work, check=True)
        out = HERE / f"chronique_{name}.docx"
        shutil.copy(work / f"_build_{name}.docx", out)
        postprocess(out)
        print("->", out.name)
    # the annex: an executable document (knitr) that sources oracle.R; no page break or heading fix-ups
    subprocess.run(["quarto", "render", "Annex.qmd", "--to", "docx"], cwd=work, check=True)
    out = HERE / "Annex.docx"
    shutil.copy(work / "Annex.docx", out)
    # Table A2 (the 54-word table) is narrower than the page: center it
    d = Document(str(out))
    for tbl in d.tables:
        if tbl.rows[0].cells[0].text.strip() == "Band":
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    d.save(str(out))
    print("-> Annex.docx")


# ---- journal limits ------------------------------------------------------------------------
def nospace(s):
    return len(re.sub(r"\s", "", s))


def block(label):
    """Text of the 'Front matter' paragraph that follows the given label paragraph."""
    m = re.search(r'Front matter"\}\n' + re.escape(label) + r'.*?\n:::\n\n::: \{custom-style="Front matter"\}\n(.*?)\n:::',
                  text, re.S)
    return re.sub(r"&nbsp;", " ", m.group(1)) if m else ""


body = text.split("# Introduction", 1)[1].split("<!--BIBLIOGRAPHY-->")[0]
body = re.sub(r"<!--.*?-->", "", body, flags=re.S)
print(f"article body: {nospace(body):,} characters without spaces (limit 50,000)")
print(f"French abstract:  {nospace(block('Résumé')):,} characters without spaces (limit 1,200)")
print(f"English abstract: {nospace(block('Abstract')):,} characters without spaces (limit 1,200)")

# ---- optional: PDF through Word ------------------------------------------------------------
if "--pdf" in sys.argv:
    for stem in [f"chronique_{name}" for name in variants] + ["Annex"]:
        docx = HERE / f"{stem}.docx"
        pdf = HERE / f"{stem}.pdf"
        ps = (
            "$w=New-Object -ComObject Word.Application; $w.Visible=$false;"
            f"$d=$w.Documents.Open('{docx}'); $d.SaveAs2('{pdf}',17); $d.Close(0); $w.Quit()"
        )
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
        print("->", pdf.name)
